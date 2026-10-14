"""Tests for the Excel report and the profile lookup that fills it."""

from openpyxl import load_workbook

from smart_investor.reporting import enrichment
from smart_investor.reporting.enrichment import Profile, fetch_profile, yfinance_symbol
from smart_investor.reporting.excel import build_workbook


class TestYfinanceSymbol:
    """Korean tickers resolve only under their market's suffix."""

    def test_us_ticker_unchanged(self):
        assert yfinance_symbol("WES", "United States", "NYQ") == "WES"

    def test_kospi_gets_ks(self):
        assert yfinance_symbol("003780", "South Korea", "KOSPI") == "003780.KS"

    def test_kosdaq_gets_kq(self):
        assert yfinance_symbol("034950", "South Korea", "KOSDAQ") == "034950.KQ"

    def test_unknown_korean_exchange_defaults_to_ks(self):
        assert yfinance_symbol("005930", "South Korea", None) == "005930.KS"


class TestPeRatio:
    """P/E comes from the trailing figure, or from market cap over earnings."""

    def test_uses_trailing_pe_when_present(self):
        assert enrichment._pe_ratio({"trailingPE": 15.481011}) == 15.48

    def test_derives_from_market_cap_and_income(self):
        """Korean tickers report no trailingPE but do report both terms."""
        info = {"marketCap": 459_023_613_952, "netIncomeToCommon": 25_583_689_728}
        assert enrichment._pe_ratio(info) == 17.94

    def test_trailing_pe_wins_over_derivation(self):
        info = {"trailingPE": 10.0, "marketCap": 100, "netIncomeToCommon": 2}
        assert enrichment._pe_ratio(info) == 10.0

    def test_loss_making_company_has_no_pe(self):
        """A negative P/E would sort as if it were cheap."""
        info = {"marketCap": 1_000_000, "netIncomeToCommon": -50_000}
        assert enrichment._pe_ratio(info) is None

    def test_zero_earnings_has_no_pe(self):
        assert enrichment._pe_ratio({"marketCap": 1_000_000, "netIncomeToCommon": 0}) is None

    def test_missing_fields_return_none(self):
        assert enrichment._pe_ratio({}) is None

    def test_negative_trailing_pe_falls_through(self):
        info = {"trailingPE": -5.0, "marketCap": 100, "netIncomeToCommon": 10}
        assert enrichment._pe_ratio(info) == 10.0


class TestFieldCleaning:
    """Blank and non-positive values become None rather than placeholders."""

    def test_clean_strips(self):
        assert enrichment._clean("  Energy  ") == "Energy"

    def test_clean_empty_is_none(self):
        assert enrichment._clean("") is None
        assert enrichment._clean("   ") is None
        assert enrichment._clean(None) is None

    def test_positive_rejects_zero_and_negative(self):
        assert enrichment._positive(0) is None
        assert enrichment._positive(-1) is None
        assert enrichment._positive(None) is None
        assert enrichment._positive("abc") is None

    def test_positive_accepts_numbers(self):
        assert enrichment._positive(1234) == 1234.0
        assert enrichment._positive("5.5") == 5.5


class TestFetchProfile:
    """Profiles degrade to empty rather than raising."""

    def _patch(self, monkeypatch, info):
        class StubTicker:
            def __init__(self, symbol):
                self.symbol = symbol

            @property
            def info(self):
                if isinstance(info, Exception):
                    raise info
                return info

        monkeypatch.setattr(enrichment.yf, "Ticker", StubTicker)

    def test_populated_profile(self, monkeypatch):
        self._patch(monkeypatch, {
            "sector": "Energy", "industry": "Oil & Gas Midstream",
            "marketCap": 20_212_500_480, "trailingPE": 15.481011,
        })

        profile = fetch_profile("WES", "United States")

        assert profile.sector == "Energy"
        assert profile.industry == "Oil & Gas Midstream"
        assert profile.market_cap == 20_212_500_480
        assert profile.pe_ratio == 15.48

    def test_empty_info_gives_blank_profile(self, monkeypatch):
        self._patch(monkeypatch, {})
        profile = fetch_profile("020710", "South Korea", "KOSDAQ")

        assert profile.ticker == "020710"
        assert profile.sector is None
        assert profile.market_cap is None

    def test_lookup_failure_is_swallowed(self, monkeypatch):
        """One unreachable ticker must not abort a whole report."""
        self._patch(monkeypatch, RuntimeError("network down"))

        assert fetch_profile("WES", "United States") == Profile(ticker="WES")


class TestBuildWorkbook:
    """The workbook keeps every requested column and explains itself."""

    def _rows(self):
        return [{
            "no": 1, "ticker": "WES", "company": "Western Midstream",
            "sector": "Energy", "industry": "Oil & Gas Midstream",
            "country": "United States", "market_cap": 20_212_500_480,
            "pe_ratio": 15.48, "price": 48.92, "change": 0.013,
            "volume": 1_785_600, "dividend_yield": 7.52,
            "years_of_dividend": 14, "debt_ratio": 58.07, "quick_ratio": 127.67,
        }]

    def test_headers_match_requested_layout(self, tmp_path):
        path = build_workbook(self._rows(), tmp_path / "r.xlsx")
        headers = [c.value for c in load_workbook(path).active[4]]

        assert headers[:11] == [
            "No.", "Ticker", "Company", "Sector", "Industry", "Country",
            "Market Cap", "P/E", "Price", "Change", "Volume",
        ]

    def test_screening_metrics_follow_the_descriptive_columns(self, tmp_path):
        path = build_workbook(self._rows(), tmp_path / "r.xlsx")
        headers = [c.value for c in load_workbook(path).active[4]]

        assert headers[11:] == ["Div Yield", "Div Years", "Debt Ratio", "Quick Ratio"]

    def test_values_written(self, tmp_path):
        path = build_workbook(self._rows(), tmp_path / "r.xlsx")
        row = [c.value for c in load_workbook(path).active[5]]

        assert row[1] == "WES"
        assert row[6] == 20_212_500_480
        assert row[9] == 0.013

    def test_missing_values_become_dashes(self, tmp_path):
        """A 0 would read as a measured value and sort with the real ones."""
        rows = self._rows()
        rows[0]["pe_ratio"] = None
        rows[0]["industry"] = None

        path = build_workbook(rows, tmp_path / "r.xlsx")
        row = [c.value for c in load_workbook(path).active[5]]

        assert row[4] == "-"
        assert row[7] == "-"

    def test_criteria_sheet_records_thresholds(self, tmp_path):
        """Tickers without their criteria are not reproducible."""
        criteria = ["dividend_yield: 배당수익률 >= 4.0%"]
        path = build_workbook(
            self._rows(), tmp_path / "r.xlsx", market="US",
            total_analyzed=7062, criteria=criteria,
        )

        sheet = load_workbook(path)["Criteria"]
        text = "\n".join(
            str(c.value) for row in sheet.iter_rows() for c in row if c.value is not None
        )
        assert criteria[0] in text
        assert "7,062" in text

    def test_header_is_frozen_and_filterable(self, tmp_path):
        sheet = load_workbook(build_workbook(self._rows(), tmp_path / "r.xlsx")).active

        assert sheet.freeze_panes == "A5"
        assert sheet.auto_filter.ref is not None

    def test_empty_rows_still_produce_a_file(self, tmp_path):
        path = build_workbook([], tmp_path / "empty.xlsx", market="KR")

        assert path.exists()
        assert load_workbook(path).active.auto_filter.ref is None

    def test_creates_missing_directories(self, tmp_path):
        path = build_workbook(self._rows(), tmp_path / "nested" / "dir" / "r.xlsx")

        assert path.exists()


class TestStrategyColumns:
    """Each strategy shows the metrics its own criteria were checked against."""

    def test_dividend_columns(self):
        from smart_investor.reporting.excel import columns_for

        headers = [c[0] for c in columns_for("dividend")]
        assert headers[11:] == ["Div Yield", "Div Years", "Debt Ratio", "Quick Ratio"]

    def test_value_columns(self):
        from smart_investor.reporting.excel import columns_for

        headers = [c[0] for c in columns_for("value")]
        assert headers[10:] == ["PER", "PBR", "ROE", "Op Margin", "Debt Ratio"]

    def test_value_replaces_market_pe_with_the_screened_one(self):
        """Two near-identical PE columns would be noise, so PER supersedes P/E."""
        from smart_investor.reporting.excel import columns_for

        keys = {c[0]: c[1] for c in columns_for("value")}
        assert "P/E" not in keys
        assert keys["PER"] == "per"

    def test_dividend_keeps_the_market_pe(self):
        from smart_investor.reporting.excel import columns_for

        keys = {c[0]: c[1] for c in columns_for("dividend")}
        assert keys["P/E"] == "pe_ratio"
        assert "PER" not in keys

    def test_unknown_strategy_falls_back_to_dividend(self):
        from smart_investor.reporting.excel import columns_for

        assert columns_for("nonsense") == columns_for("dividend")


class TestMultiSheetWorkbook:
    """Both strategies share one workbook so they open together."""

    def _sheets(self):
        return [
            {
                "strategy": "dividend",
                "rows": [{"no": 1, "ticker": "020710", "dividend_yield": 9.26}],
                "criteria": ["dividend_yield: 배당수익률 >= 4.0%"],
                "total_analyzed": 2763,
            },
            {
                "strategy": "value",
                "rows": [{"no": 1, "ticker": "088130", "pbr": 0.41, "roe": 39.81}],
                "criteria": ["per: PER <= 15.0"],
                "total_analyzed": 2763,
            },
        ]

    def test_one_sheet_per_strategy_plus_criteria(self, tmp_path):
        from smart_investor.reporting.excel import build_multi_workbook

        path = build_multi_workbook(self._sheets(), tmp_path / "both.xlsx", market="KR")

        assert load_workbook(path).sheetnames == ["배당주 KR", "가치주 KR", "Criteria"]

    def test_each_sheet_uses_its_own_columns(self, tmp_path):
        from smart_investor.reporting.excel import build_multi_workbook

        workbook = load_workbook(
            build_multi_workbook(self._sheets(), tmp_path / "both.xlsx", market="KR")
        )

        assert [c.value for c in workbook["배당주 KR"][4]][11] == "Div Yield"
        assert [c.value for c in workbook["가치주 KR"][4]][10] == "PER"

    def test_criteria_sheet_covers_every_strategy(self, tmp_path):
        from smart_investor.reporting.excel import build_multi_workbook

        sheet = load_workbook(
            build_multi_workbook(self._sheets(), tmp_path / "both.xlsx", market="KR")
        )["Criteria"]
        text = "\n".join(
            str(c.value) for row in sheet.iter_rows() for c in row if c.value is not None
        )

        assert "[배당주]" in text and "[가치주]" in text
        assert "per: PER <= 15.0" in text
        assert "dividend_yield: 배당수익률 >= 4.0%" in text

    def test_empty_strategy_still_gets_a_sheet(self, tmp_path):
        """A strategy with no matches must not vanish from the report."""
        from smart_investor.reporting.excel import build_multi_workbook

        sheets = self._sheets()
        sheets[1]["rows"] = []

        workbook = load_workbook(
            build_multi_workbook(sheets, tmp_path / "both.xlsx", market="US")
        )
        assert "가치주 US" in workbook.sheetnames
